import os
import json
import re
import zipfile
import io
from PIL import Image
import tempfile

# 尝试使用python-docx库获取样式信息
try:
    from docx import Document
    docx_available = True
except ImportError:
    docx_available = False
    print("python-docx库未安装，无法使用样式信息提取功能。")

# 尝试导入unstructured库的必要组件
try:
    from unstructured.partition.docx import partition_docx
    # 尝试导入Heading和Text，如果失败则不使用
    try:
        from unstructured.documents.elements import Heading, Text
    except ImportError:
        print("unstructured库中未找到Heading或Text类，将使用替代方式处理元素。")
        Heading = None
        Text = None
    unstructured_available = True
except ImportError:
    unstructured_available = False
    print("unstructured库未安装，无法使用回退功能。")

docs_path = "d:\Desktop\coding\DocumentModule\DocModule\docx\合同管理.docx"


class DocumentStructureExtractor:
    def __init__(self, docx_path):
        """初始化文档结构提取器"""
        self.docx_path = docx_path
        self.structure = {
            "type": "document",
            "cover": [],  # 封皮部分
            "toc": [],    # 目录部分
            "content": [] # 正文部分
        }
        self.current_section = "cover"  # 默认从封皮开始
        self.current_headings = []
        self.element_counter = 0  # 统一计数器，用于index
        self.in_table_of_contents = False  # 是否在目录中
        self.toc_started = False  # 目录是否已开始
        self.content_started = False  # 正文是否已开始

    def extract_structure(self):
        """提取文档结构并返回JSON格式"""
        if docx_available:
            print("使用python-docx库提取文档结构...")
            return self._extract_with_docx()
        elif unstructured_available:
            print("使用unstructured库提取文档结构...")
            # 解析文档
            print(f"开始解析文档: {self.docx_path}")
            self.elements = partition_docx(filename=self.docx_path)
            print(f"共提取 {len(self.elements)} 个元素")

            # 处理每个元素
            for i, element in enumerate(self.elements, 1):
                print(f"处理元素 {i}/{len(self.elements)}: {element.__class__.__name__}")
                self._process_element(element)

            return self.structure
        else:
            print("错误: 未找到可用的文档处理库。")
            return self.structure

    def _extract_with_docx(self):
        """使用python-docx库提取文档结构"""
        doc = Document(self.docx_path)
        
        # 获取所有元素（段落、表格和图片）并按文档中的实际顺序处理
        elements = []
        
        # 用zip方式打开docx，方便读取图片
        with zipfile.ZipFile(self.docx_path) as zf:
            # 遍历文档的XML元素，以获取段落、表格和图片的实际顺序
            for child in doc.element.body.iterchildren():
                tag_name = child.tag.split('}')[-1]  # 获取不带命名空间的标签名
                if tag_name == 'p':  # 段落
                    # 找到对应的paragraph对象
                    for para in doc.paragraphs:
                        if para._element == child:
                            elements.append(('paragraph', para))
                            break
                elif tag_name == 'tbl':  # 表格
                    # 找到对应的table对象
                    for table in doc.tables:
                        if table._element == child:
                            elements.append(('table', table))
                            break
                # 检查段落中的图片
                if tag_name == 'p':
                    for para in doc.paragraphs:
                        if para._element == child:
                            # 检查段落中是否包含图片
                            for drawing in para._element.xpath('.//w:drawing'):
                                blip = drawing.find('.//a:blip',
                                                    namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'})
                                if blip is None:
                                    continue
                                rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                                if rid is None:
                                    continue
                                rel = doc.part.rels.get(rid)
                                if rel is None or not rel.target_ref.startswith('media/'):
                                    continue
                                try:
                                    img_data = zf.read(f'word/{rel.target_ref}')
                                    w, h = Image.open(io.BytesIO(img_data)).size
                                    if w > 100 or h > 100:  # 过滤小图片
                                        elements.append(('image', {'para': para, 'rel': rel, 'img_data': img_data, 'w': w, 'h': h}))
                                except Exception:
                                    continue
                            break
        
        # 按文档中的实际顺序处理元素
        for element_type, element in elements:
            if element_type == 'paragraph':
                para = element
                if not para.text.strip():
                    continue
                # 获取段落样式名称
                style_name = para.style.name
                # print(f"段落文本: '{para.text.strip()}', 样式名称: '{style_name}'")#调试用输出

                # 确定当前段落属于哪个部分
                self._determine_section(para.text.strip(), style_name)

                # 处理标题
                # 匹配多种标题样式格式，如'Sunline标题4'或'Heading 1'
                level_match = re.search(r'(标题|Heading)\s*(\d+)', style_name)
                if level_match:
                    self.element_counter += 1
                    level = int(level_match.group(2))
                    heading_node = {
                        "type": "heading", 
                        "index": self.element_counter, 
                        "level": level, 
                        "text": para.text.strip(), 
                        "children": []
                    }
                    self._add_heading(heading_node)
                else:
                    # 处理目录部分的层级结构
                    if self.current_section == "toc":
                        # 匹配目录项的编号格式 (如: 1, 1.1, 2, 2.1.1)
                        toc_level_match = re.match(r'^(\d+\.)*\d+\s+', para.text.strip())
                        if toc_level_match:
                            self.element_counter += 1
                            # 计算层级 (根据点的数量+1)
                            level = len(toc_level_match.group(0).split('.'))
                            # 提取标题文本 (去除编号部分)
                            toc_text = para.text.strip()[len(toc_level_match.group(0)):]
                            toc_heading_node = {
                                "type": "heading", 
                                "index": self.element_counter, 
                                "level": level, 
                                "text": toc_text, 
                                "children": []
                            }
                            # 目录层级管理
                            while self.toc_current_headings and self.toc_current_headings[-1]['level'] >= level:
                                self.toc_current_headings.pop()

                            if self.toc_current_headings:
                                self.toc_current_headings[-1]['children'].append(toc_heading_node)
                            else:
                                self.structure['toc'].append(toc_heading_node)

                            self.toc_current_headings.append(toc_heading_node)
                            continue

                    self.element_counter += 1
                    paragraph_node = {
                        "type": "paragraph", 
                        "index": self.element_counter, 
                        "text": para.text.strip()
                    }
                    self._add_paragraph(paragraph_node)
            elif element_type == 'table':
                table = element
                self.element_counter += 1
                table_data = []
                # 遍历表格的行
                for row in table.rows:
                    row_data = []
                    # 遍历行中的单元格
                    for cell in row.cells:
                        row_data.append(cell.text.strip())
                    table_data.append(row_data)
                # 创建表格节点
                table_node = {
                    "type": "table", 
                    "index": self.element_counter, 
                    "row": len(table_data),
                    "col": len(table_data[0]) if table_data else 0,
                    "data": table_data
                }
                # 添加表格到结构中
                self._add_table(table_node)
                print(f"提取表格: {len(table.rows)}行, {len(table.columns)}列")
            elif element_type == 'image':
                # 处理图片元素
                image_info = element
                para = image_info['para']
                rel = image_info['rel']
                img_data = image_info['img_data']
                w = image_info['w']
                h = image_info['h']
                
                # 确定当前段落属于哪个部分
                style_name = para.style.name if hasattr(para, 'style') else ''
                self._determine_section(para.text.strip(), style_name)
                
                self.element_counter += 1
                original_name = rel.target_ref.split('/')[-1]
                file_size = len(img_data)
                
                # 创建图片节点
                image_node = {
                    "type": "image", 
                    "index": self.element_counter, 
                    "original_name": original_name,
                    "dimensions": f"{w}x{h}",
                    "file_size": file_size
                }
                # 添加图片到结构中
                self._add_image(image_node)
                print(f"提取图片: {original_name}, 尺寸: {w}x{h}, 大小: {file_size}字节")

        return self.structure

    def _determine_section(self, text, style_name):
        """确定当前段落属于哪个部分（封皮、目录或正文）"""
        # 检查是否是目录开始
        if not self.toc_started and (text.strip() == "目录" or style_name == "目录"):
            self.current_section = "toc"
            self.toc_started = True
            self.toc_current_headings = []  # 初始化目录标题层级管理
            print("进入目录部分")
            return

        # 检查是否是正文开始
        if self.toc_started and not self.content_started:
            # 假设目录后的第一个标题是正文开始
            if re.search(r'(标题|Heading)\s*1', style_name):
                self.current_section = "content"
                self.content_started = True
                print("进入正文部分")
                return

    def _add_heading(self, heading_node):
        """添加标题到文档结构中"""
        # 实现标题添加逻辑
        level = heading_node['level']
        # 找到合适的位置添加标题
        while self.current_headings and self.current_headings[-1]['level'] >= level:
            self.current_headings.pop()

        # 根据当前部分添加到相应位置
        if self.current_section == "content":
            if self.current_headings:
                self.current_headings[-1]['children'].append(heading_node)
            else:
                self.structure['content'].append(heading_node)

            self.current_headings.append(heading_node)
        elif self.current_section == "toc":
            # 目录部分的标题直接添加到toc数组
            self.structure['toc'].append(heading_node)
        else:
            # 封皮部分的标题直接添加到cover数组
            self.structure['cover'].append(heading_node)

    def _add_paragraph(self, paragraph_node):
        """添加段落到文档结构中"""
        # 根据当前部分添加到相应位置
        if self.current_section == "content":
            if self.current_headings:
                self.current_headings[-1]['children'].append(paragraph_node)
            else:
                self.structure['content'].append(paragraph_node)
        elif self.current_section == "toc":
            # 目录部分的段落直接添加到toc数组
            self.structure['toc'].append(paragraph_node)
        else:
            # 封皮部分的段落直接添加到cover数组
            self.structure['cover'].append(paragraph_node)

    def _add_table(self, table_node):
        """添加表格到文档结构中"""
        # 根据当前部分添加到相应位置
        if self.current_section == "content":
            if self.current_headings:
                self.current_headings[-1]['children'].append(table_node)
            else:
                self.structure['content'].append(table_node)
        elif self.current_section == "toc":
            self.structure['toc'].append(table_node)
        else:
            self.structure['cover'].append(table_node)

    def _add_image(self, image_node):
        """添加图片到文档结构中"""
        # 根据当前部分添加到相应位置
        if self.current_section == "content":
            if self.current_headings:
                self.current_headings[-1]['children'].append(image_node)
            else:
                self.structure['content'].append(image_node)
        elif self.current_section == "toc":
            self.structure['toc'].append(image_node)
        else:
            self.structure['cover'].append(image_node)

    def _process_element(self, element):
        """处理单个元素，根据类型添加到结构中"""
        element_type = element.__class__.__name__
        text = getattr(element, 'text', '').strip()

        # 过滤空文本
        if not text.strip() and element_type != 'Table':
            return

        # 直接处理所有元素
        heading_processed = self._process_heading(element)
        if not heading_processed:
            # 如果未作为标题处理，则根据元素类型处理
            if element_type == 'Paragraph':
                self._process_paragraph(element)
            elif element_type == 'Image':
                self._process_image(element)
            elif element_type == 'Table':
                self._process_table(element)
            else:
                # 对于其他类型，尝试作为段落处理
                if text:
                    paragraph = type('Paragraph', (object,), {'text': text})
                    self._process_paragraph(paragraph)

    def _process_heading(self, heading):
        """处理标题元素，更新当前标题层级"""
        text = getattr(heading, 'text', '').strip()
        if not text:
            return False

        # 放宽标题文本长度限制
        if len(text) > 150:
            # 作为段落处理
            self.paragraph_counter += 1
            paragraph_id = f"p{self.paragraph_counter}"
            paragraph_node = {
                "type": "paragraph",
                "id": paragraph_id,
                "text": text
            }
            if self.current_headings:
                self.current_headings[-1]["children"].append(paragraph_node)
            else:
                self.structure["children"].append(paragraph_node)
            return False

        # 更灵活的标题级别判断
        level = None

        # 1. 检查样式名
        style_name = None
        # 尝试多种方式获取样式信息
        print(f"检查heading对象类型: {type(heading)}")
        print(f"heading属性列表: {dir(heading)}")
        
        if hasattr(heading, 'style'):
            print(f"style属性类型: {type(heading.style)}")
            print(f"style属性内容: {heading.style}")
            if hasattr(heading.style, 'name'):
                style_name = heading.style.name
                print(f"通过heading.style.name获取样式名: {style_name}")
            else:
                style_name = heading.style
                print(f"通过heading.style获取样式名: {style_name}")
        elif hasattr(heading, 'formatting'):
            print(f"formatting属性类型: {type(heading.formatting)}")
            print(f"formatting属性内容: {heading.formatting}")
            if hasattr(heading.formatting, 'font'):
                print(f"font属性类型: {type(heading.formatting.font)}")
                print(f"font属性内容: {heading.formatting.font}")
                if hasattr(heading.formatting.font, 'name'):
                    style_name = heading.formatting.font.name
                    print(f"通过heading.formatting.font.name获取样式名: {style_name}")
        elif hasattr(heading, 'attributes'):
            print(f"attributes属性内容: {heading.attributes}")
            if 'style_name' in heading.attributes:
                style_name = heading.attributes['style_name']
                print(f"通过attributes获取样式名: {style_name}")
            elif 'style' in heading.attributes:
                style_name = heading.attributes['style']
                print(f"通过attributes获取样式名: {style_name}")
        
        print(f"当前处理的文本: '{heading.text}'")
        print(f"最终样式名称: {style_name}")
        
        level = None
        if style_name:
            # 检查样式名中是否包含标题相关关键词
            if '标题' in style_name:
                # 尝试从样式名中提取级别数字
                level_match = re.search(r'(\d+)', style_name)
                if level_match:
                    level = int(level_match.group(1))
                    print(f"通过样式名'{style_name}'识别标题级别: {level}")
                else:
                    level = 1  # 默认级别
                    print(f"样式名包含'标题'，但未找到级别，默认级别: {level}")
        
        if level is None:
            print("未找到样式信息，尝试其他方法识别标题级别")
            print(f"检查heading对象类型: {type(heading)}")
            print(f"heading属性列表: {dir(heading)}")
            
            # 检查heading本身是否有样式相关属性
            style_name = None
            print(f"当前处理的文本: '{heading.text}'")
            if hasattr(heading, 'style'):
                style_name = heading.style
            elif hasattr(heading, 'formatting'):
                style_name = heading.formatting
            elif hasattr(heading, 'font'):
                style_name = heading.font
            elif hasattr(heading, 'attributes') and 'style' in heading.attributes:
                style_name = heading.attributes['style']
            
            if style_name:
                print(f"找到样式信息: '{style_name}'")
                if '标题' in str(style_name):
                    # 提取数字作为标题级别
                    level_match = re.search(r'\d+', str(style_name))
                    if level_match:
                        level = int(level_match.group())
                        print(f"通过样式名'{style_name}'识别标题级别: {level}")
            else:
                print("未找到样式信息")

        # 3. 基于数字格式识别标题层级
        if level is None:
            print(f"尝试通过数字格式识别标题: '{heading.text.strip()}'")
            # 更精确的数字格式匹配，支持如"1 标题"、"1.1 标题"、"1.1.1 标题"等格式
            numbering_match = re.match(r'^(\d+(\.\d+)*)\s+', heading.text.strip())
            
            if numbering_match:
                # 提取数字层级
                numbering = numbering_match.group(1)
                level = len(numbering.split('.'))
                print(f"通过数字格式'{numbering}'识别标题级别: {level}")
            else:
                # 检查是否是章节标题关键词
                chapter_keywords = ['第[一二三四五六七八九十]章', '第[一二三四五六七八九十]节', '目录', '概述', '引言', '结论', '参考文献']
                for keyword in chapter_keywords:
                    if re.match(f'^{keyword}', heading.text.strip()):
                        level = 1
                        print(f"通过章节关键词'{heading.text.strip()}'识别标题级别: {level}")
                        break
                else:
                    # 检查是否是无编号的顶级标题
                    if re.match(r'^[\u4e00-\u9fa5]+$', heading.text.strip()) and len(heading.text.strip()) <= 10:
                        # 检查是否是已知的子标题关键词
                        subheading_keywords = ['目的', '目标', '读者', '建议', '术语', '缩写词', '描述', '综述', '列表', '规则']
                        is_subheading = False
                        for keyword in subheading_keywords:
                            if keyword in heading.text.strip():
                                is_subheading = True
                                level = 2
                                print(f"通过子标题关键词'{keyword}'识别标题级别: {level}")
                                break
                        if not is_subheading:
                            level = 1
                            print(f"通过中文标题识别为顶级标题: '{heading.text.strip()}'")
                    else:
                        # 默认级别
                        level = None
                        print(f"无法识别标题级别: '{heading.text.strip()}'")
        # 2. 检查是否是Heading类型元素
        if level is None:
            element_type = heading.__class__.__name__
            if element_type == 'Heading':
                if hasattr(heading, 'level'):
                    level = heading.level
                else:
                    level = 1

        # 3. 如果不是Heading类型或无法获取级别，则通过格式判断
        if level is None:
            # 检查是否是目录中的标题项 (如 "1 概述", "2.1 功能综述")
            if re.match(r'^\d+\s', text):
                level = 1
            elif re.match(r'^\d+\.\d+\s', text):
                level = 2
            elif re.match(r'^\d+\.\d+\.\d+\s', text):
                level = 3
            elif re.match(r'^\d+\.\d+\.\d+\.\d+\s', text):
                level = 4
            # 检查中文编号格式
            elif text.startswith(('一、', '二、', '三、', '四、', '五、', '六、', '七、', '八、', '九、', '十、', '十一、', '十二、')):
                level = 1
            elif text.startswith(('（一）', '（二）', '（三）', '（四）', '（五）', '（六）', '（七）', '（八）', '（九）', '（十）')):
                level = 2
            elif text.startswith(('1)', '2)', '3)', '4)', '5)', '6)', '7)', '8)', '9)', '10)')) or text.startswith(('①', '②', '③', '④', '⑤', '⑥', '⑦', '⑧', '⑨', '⑩')):
                level = 3
            elif text.startswith(('a.', 'b.', 'c.', 'd.', 'e.', 'f.', 'g.', 'h.', 'i.', 'j.')) or text.startswith(('A.', 'B.', 'C.', 'D.', 'E.', 'F.', 'G.', 'H.', 'I.', 'J.')):
                level = 4
            # 检查是否是粗体或大号字体的文本
            elif hasattr(heading, 'metadata') and hasattr(heading.metadata, 'font') and 'bold' in heading.metadata.font.lower():
                level = 1
            elif hasattr(heading, 'metadata') and hasattr(heading.metadata, 'font_size') and heading.metadata.font_size and float(heading.metadata.font_size) > 12:
                level = 1
            # 检查样式是否包含heading
            elif hasattr(heading, 'metadata') and hasattr(heading.metadata, 'style') and 'heading' in heading.metadata.style.lower():
                level = 1

        # 4. 如果仍然没有级别，尝试根据上下文判断
        if level is None:
            # 检查是否是常见的章节标题词汇
            chapter_keywords = ['概述', '引言', '背景', '目的', '范围', '定义', '功能描述', '系统架构', '实现方案', '测试计划', '结论', '参考文献', '本文档的目的', '目标读者和阅读建议', '术语和缩写词', '共通描述']
            for keyword in chapter_keywords:
                if keyword in text and len(text) <= 30:
                    if keyword in ['本文档的目的', '目标读者和阅读建议', '术语和缩写词', '共通描述']:
                        level = 2
                    else:
                        level = 1
                    break

        # 如果没有匹配到任何标题格式，作为段落处理
        if level is None:
            self.paragraph_counter += 1
            paragraph_id = f"p{self.paragraph_counter}"
            paragraph_node = {
                "type": "paragraph",
                "id": paragraph_id,
                "text": text
            }
            if self.current_headings:
                self.current_headings[-1]["children"].append(paragraph_node)
            else:
                self.structure["children"].append(paragraph_node)
            return False

        print(f"找到标题: 级别 {level}, 内容: {text[:30]}...")

        # 清理标题文本，移除可能的页码和格式字符
        clean_text = re.sub(r'\t\d+$', '', text)  # 移除尾部的制表符和数字
        clean_text = re.sub(r'^\d+\.\s*', '', clean_text)  # 移除开头的数字和点
        clean_text = clean_text.strip()

        # 检查当前层级是否在已有的标题层级中
        while self.current_headings and self.current_headings[-1]["level"] >= level:
            self.current_headings.pop()

        # 创建新的标题节点
        self.heading_counter += 1
        heading_id = f"h{self.heading_counter}"
        heading_node = {
            "type": "heading",
            "id": heading_id,
            "level": level,
            "text": clean_text,
            "children": []
        }

        # 添加标题节点到结构中
        if self.current_headings:
            self.current_headings[-1]["children"].append(heading_node)
        else:
            self.structure["children"].append(heading_node)

        # 更新当前标题层级
        self.current_headings.append(heading_node)
        return True

    def _process_paragraph(self, paragraph):
        """处理段落元素，为段落编号并添加到当前标题层级下"""
        text = paragraph.text.strip()
        if not text:  # 跳过空段落
            return

        self.paragraph_counter += 1
        paragraph_id = f"p{self.paragraph_counter}"

        paragraph_node = {
            "type": "paragraph",
            "id": paragraph_id,
            "text": text
        }

        # 不需要单独的paragraphs列表，直接添加到树结构中

        # 添加到当前标题层级下
        if self.current_headings:
            self.current_headings[-1]["children"].append(paragraph_node)
        else:
            self.structure["children"].append(paragraph_node)

    def _process_image(self, image):
        """处理图片元素，存储图片名称"""
        # 尝试获取图片名称，如果没有则使用默认名称
        image_name = getattr(image.metadata, 'filename', f"image_{len(self.structure['children']) + 1}")

        image_node = {
            "type": "image",
            "name": image_name,
            "metadata": {
                "height": getattr(image.metadata, 'height', None),
                "width": getattr(image.metadata, 'width', None)
            }
        }

        # 添加到当前标题层级下
        if self.current_headings:
            self.current_headings[-1]["children"].append(image_node)
        else:
            self.structure["children"].append(image_node)

    def _process_table(self, table):
        """处理表格元素"""
        # 尝试从不同途径获取表格行列信息
        num_rows = None
        num_columns = None
        try:
            # 方法1: 尝试从metadata获取
            num_rows = getattr(table.metadata, 'row_count', None)
            num_columns = getattr(table.metadata, 'col_count', None)
            
            # 方法2: 如果metadata没有，尝试直接访问rows属性
            if num_rows is None and hasattr(table, 'rows'):
                num_rows = len(table.rows)
                if num_rows > 0 and hasattr(table.rows[0], 'cells'):
                    num_columns = len(table.rows[0].cells)
                    
            # 方法3: 尝试通过__dict__访问
            if num_rows is None and 'rows' in table.__dict__:
                num_rows = len(table.__dict__['rows'])
                if num_rows > 0 and 'cells' in table.__dict__['rows'][0].__dict__:
                    num_columns = len(table.__dict__['rows'][0].__dict__['cells'])
        except Exception as e:
            print(f"提取表格行列信息出错: {e}")
            
        table_node = {
            "type": "table",
            "text": str(table),  # 简化处理，仅存储表格文本表示
            "num_rows": num_rows,
            "num_columns": num_columns
        }

        # 添加到当前标题层级下
        if self.current_headings:
            self.current_headings[-1]["children"].append(table_node)
        else:
            self.structure["children"].append(table_node)

    def save_to_json(self, output_path):
        """将提取的文档结构保存为JSON文件"""
        # 后处理：去重
        def remove_duplicates(node):
            if not isinstance(node, dict):
                return node

            # 处理子节点
            if "children" in node:
                seen = {}
                unique_children = []
                for child in node["children"]:
                    # 创建一个基于类型和文本的唯一键
                    if child.get("type") == "paragraph" or child.get("type") == "heading":
                        # 对于段落和标题，使用类型+文本作为键
                        key = f"{child.get('type')}:{child.get('text', '').strip()}"
                    elif child.get("type") == "table":
                        # 对于表格，使用类型+文本摘要作为键
                        text_summary = child.get('text', '')[:100].strip()
                        key = f"{child.get('type')}:{text_summary}"
                    else:
                        # 对于其他类型，使用类型+id作为键
                        key = f"{child.get('type')}:{child.get('id', '')}"

                    if key not in seen:
                        seen[key] = True
                        # 递归处理子节点
                        if "children" in child:
                            child["children"] = remove_duplicates({"children": child["children"]})["children"]
                        unique_children.append(child)
                node["children"] = unique_children

            return node

        # 应用去重
        self.structure = remove_duplicates(self.structure)

        # 保存到JSON文件
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(self.structure, f, ensure_ascii=False, indent=2)


# 主函数
def extract_document(input_source, save_to_file=True, return_json=True):
    # 判断输入类型并获取文件路径
    temp_created = False
    if hasattr(input_source, 'read'):  # 检查是否为file对象
        # 创建临时文件
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as temp_file:
            temp_file.write(input_source.read())
            docs_path = temp_file.name
        temp_created = True
    else:
        docs_path = input_source

    try:
        # 提取文档结构
        extractor = DocumentStructureExtractor(docs_path)
        structure = extractor.extract_structure()

        # 可选：保存到文件
        if save_to_file:
            output_dir = os.path.join(os.path.dirname(__file__), 'output')
            os.makedirs(output_dir, exist_ok=True)
            file_name = os.path.basename(docs_path)
            file_name_without_ext = os.path.splitext(file_name)[0]
            output_path = os.path.join(output_dir, f"{file_name_without_ext}_get.json")
            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(structure, f, ensure_ascii=False, indent=2)
            print(f"文档结构已保存至: {output_path}")

        # 根据return_json参数决定返回类型
        if return_json:
            return json.dumps(structure, ensure_ascii=False, indent=2)
        return structure
    finally:
        # 清理临时文件
        if temp_created:
            os.remove(docs_path)

def main():
    import argparse
    parser = argparse.ArgumentParser(description='Convert DOCX documents to JSON format')
    parser.add_argument('input_file', help='Path to the input DOCX file')
    parser.add_argument('output_file', nargs='?', help='Path to the output JSON file (optional)')
    args = parser.parse_args()

    # 如果提供了输出文件路径，则使用该路径
    if args.output_file:
        structure = extract_document(args.input_file, save_to_file=False, return_json=False)
        with open(args.output_file, 'w', encoding='utf-8') as f:
            json.dump(structure, f, ensure_ascii=False, indent=2)
        print(f"文档结构已保存至: {args.output_file}")
    else:
        # 否则使用默认行为
        extract_document(args.input_file)

if __name__ == "__main__":
    main()